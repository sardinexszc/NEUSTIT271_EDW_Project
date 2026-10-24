"""Compile the documentation into one submission report: docs/submission/DRIW_Final_Report.docx (+ .pdf).

The chapters are the Markdown files in docs/, so the report is rebuilt whenever they change.
Run the generators first when the data changes:

    python docs/build/data_dictionary.py
    python docs/build/diagrams.py
    python docs/build/screenshots.py      (XAMPP must be serving http://localhost/driw/)
    python ingestion/benchmark_queries.py
    python docs/build/build_report.py     (Word updates the table of contents and exports the PDF)
"""

import re
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import pymupdf
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT_DIR = DOCS / "submission"
DOCX = OUT_DIR / "DRIW_Final_Report.docx"

# academic layout: A4, 2.54 cm margins, Times New Roman 12 pt, 1.5 lines, justified body, black text
FONT, CODE_FONT = "Times New Roman", "Consolas"
MARGIN_CM = 2.54
TEXT_WIDTH_CM = 21 - 2 * MARGIN_CM
NAVY = RGBColor(0, 0, 0)            # headings, captions and code are black, like the body
GREY = RGBColor(0x40, 0x40, 0x40)   # caption discussions and notes
IMAGE_WIDTH_CM = TEXT_WIDTH_CM
IMAGE_MAX_HEIGHT_CM = 21.0   # taller images are split at a blank gap, not shrunk
IMAGE_FIT_HEIGHT_CM = 18.5   # pieces taller than this are narrowed so the caption fits beneath
HEADER_FILL = "E7E6E6"
CODE_FILL = "F2F2F2"
# title page: filled in when known, left out when empty
UNIVERSITY, DEPARTMENT, SUBMITTED_TO = "", "", ""

# (chapter title, [(markdown file, heading-level shift)]); front matter is unnumbered
FRONT_MATTER = [
    ("Executive Summary", [("executive-summary.md", 0)]),
    ("Requirements Coverage", [("requirements-traceability.md", 0)]),
    ("List of Abbreviations", [("abbreviations.md", 0)]),
]
CHAPTERS = [
    ("Business Case and Requirements Analysis", [("business-case.md", 0)]),
    ("Data Warehouse Design", [("warehouse-design.md", 0), ("dimensional-model.md", 1)]),
    ("Data Dictionary and Metadata", [("data-dictionary.md", 0)]),
    ("ETL Design", [("etl-design.md", 0)]),
    ("Data Quality Framework", [("data-quality-framework.md", 0)]),
    ("Data Governance, Security and Privacy", [("governance.md", 0)]),
    ("Analytics: OLAP and Dashboards", [("analytics.md", 0)]),
    ("Testing and Validation", [("validation-report.md", 0), ("query-performance.md", 1)]),
    ("Implementation and Live Demo", [("implementation.md", 0)]),
    ("Enterprise Positioning and Data Exchange", [("enterprise-positioning.md", 0)]),
    ("Conclusion and Recommendations", [("conclusion.md", 0)]),
]
BACK_MATTER = [
    ("References", [("references.md", 0)]),
]
# how a link to another file reads in print, where the label is the file name itself
IN_PRINT = {"dimensional-model.md": "the Dimensional Model section of Chapter 2",
            "query-performance.md": "the Query Performance Test Report (Chapter 8)",
            "sample-queries.sql": "Appendix A"}
IN_PRINT.update({name: f"Chapter {n}" for n, (_, files) in enumerate(CHAPTERS, start=1)
                 for name, _ in files if name not in IN_PRINT})


def split_tall(path: Path) -> list[Path]:
    """Cut a raster image taller than a page into page-sized pieces, each cut on a blank row."""
    if path.suffix.lower() != ".png":
        return [path]
    pix = pymupdf.Pixmap(str(path))
    limit = int(pix.width * IMAGE_MAX_HEIGHT_CM / IMAGE_WIDTH_CM)
    if pix.height <= limit:
        return [path]
    n, stride, data = pix.n, pix.stride, pix.samples
    bottom = (pix.height - 1) * stride
    blank = data[bottom:bottom + n] * pix.width     # a full row of the page background

    def is_blank(y: int) -> bool:
        return data[y * stride:y * stride + pix.width * n] == blank

    cuts, top = [], 0
    while pix.height - top > limit:
        y = top + limit
        while y > top + limit // 3 and not is_blank(y):
            y -= 1
        if y <= top + limit // 3:            # no gap found: cut at the limit
            y = top + limit
        cuts.append((top, y))
        top = y
    cuts.append((top, pix.height))

    out_dir = Path(tempfile.mkdtemp(prefix="driw-img-"))
    doc = pymupdf.open(str(path))
    page = doc[0]
    k = page.rect.height / pix.height
    zoom = pymupdf.Matrix(pix.width / page.rect.width, pix.width / page.rect.width)
    parts = []
    for i, (a, b) in enumerate(cuts):
        part = out_dir / f"{path.stem}_{i}.png"
        page.get_pixmap(clip=pymupdf.Rect(0, a * k, page.rect.width, b * k), matrix=zoom).save(str(part))
        parts.append(part)
    doc.close()
    return parts


# ----------------------------------------------------------------------------- docx helpers

def shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def add_field(paragraph, instruction: str,
              result: str = "Right-click and choose Update Field to build the table of contents."):
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate"),
                      ("w:t", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        elif tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = instruction
        else:
            el.text = result
        run._r.append(el)
    return run


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    # (size, space before, space after, italic)
    for level, (size, before, after, italic) in {1: (16, 18, 8, False), 2: (14, 12, 6, False),
                                                 3: (12, 10, 4, False), 4: (12, 8, 4, True)}.items():
        h = styles[f"Heading {level}"]
        h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(size), True, italic
        fonts = h.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)   # the theme font would override the named one
        h.font.color.rgb = NAVY
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.line_spacing = 1.15
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = level == 1   # every chapter starts a page
    cap = styles["Caption"]
    cap.font.name, cap.font.size, cap.font.color.rgb = FONT, Pt(10), NAVY
    cap.paragraph_format.line_spacing = 1.0


def footer(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DRIW · IT271 Enterprise Data Warehousing · page ")
    run.font.size, run.font.color.rgb = Pt(9), GREY
    add_field(p, "PAGE", "")
    p.runs[-1].font.size = Pt(9)


def title_page(doc: Document) -> None:
    """Centred, black, no decoration: institution, title, purpose, author, date."""
    def line(text, size=12, bold=False, italic=False, before=0):
        if not text:
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(4)
        r = p.add_run(text)
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic

    line(UNIVERSITY.upper(), 13, bold=True)
    line(DEPARTMENT.upper(), 12)
    line("DISASTER RESPONSE INTELLIGENCE WAREHOUSE (DRIW)", 18, bold=True, before=90)
    line("Design, Implementation, and Validation of an Enterprise Data Warehouse for Disaster Impact, "
         "Evacuation, Relief, and Emergency-Response Analytics in the Philippines, 2020–2025", 13, before=10)
    line("A Project Report Submitted in Partial Fulfillment of the Requirements for", 12, before=70)
    line("IT271 – Enterprise Data Warehousing", 12, bold=True)
    line("Submitted by", 12, italic=True, before=50)
    line("Ivan Christian Salinas", 12, bold=True)
    if SUBMITTED_TO:
        line("Submitted to", 12, italic=True, before=18)
        line(SUBMITTED_TO, 12, bold=True)
    line(date.today().strftime("%B %Y"), 12, before=50)
    line("Repository: github.com/sardinexszc/NEUSTIT271_EDW_Project", 10, before=40)
    line("Data notice: event names, dates and PSGC geography are real. All damage, evacuation, relief "
         "and response figures are synthetic, generated for demonstration.", 10, italic=True)
    # new pages start with "page break before" on their heading, never an empty break paragraph:
    # one that falls at the top of a page would leave that page blank
    for heading, instruction in (("Contents", 'TOC \\o "1-2" \\h \\z \\u'),
                                 ("List of Figures", 'TOC \\h \\z \\c "Figure"'),
                                 ("List of Tables", 'TOC \\h \\z \\c "Table"')):
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(heading.upper())
        r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, NAVY
        add_field(doc.add_paragraph(), instruction)


# ----------------------------------------------------------------------------- markdown -> docx

class Renderer:
    def __init__(self, doc: Document, base: Path):
        self.doc, self.base = doc, base
        self.md = MarkdownIt("commonmark").enable("table")
        self.section = 0
        self.counts = {"Figure": 0, "Table": 0}
        self.uncaptioned = []
        self.hanging = False   # list items as reference entries (hanging indent)

    @staticmethod
    def caption_of(children, kind: str):
        """A paragraph that opens with **Figure: Title.** or **Table: Title.** -> (title, discussion tokens)."""
        children = [c for c in children if not (c.type == "text" and c.content == "")]
        if len(children) < 3 or children[0].type != "strong_open":
            return None
        close = next((k for k, c in enumerate(children) if c.type == "strong_close"), None)
        title = "".join(c.content for c in children[1:close] if c.type in ("text", "code_inline"))
        if not title.startswith(f"{kind}:"):
            return None
        rest = children[close + 1:]
        while rest and rest[0].type in ("text", "softbreak") and not rest[0].content.strip():
            rest = rest[1:]
        return title[len(kind) + 1:].strip().rstrip("."), rest

    def caption(self, kind: str, title: str, discussion) -> None:
        """Numbered caption (a Word SEQ field, so the lists of figures and tables pick it up),
        then its short discussion in a paragraph of its own so the lists show the title only."""
        self.counts[kind] += 1
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "Figure" else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4 if kind == "Figure" else 8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = kind == "Table" or bool(discussion)
        runs = [p.add_run(f"{kind} "), add_field(p, f"SEQ {kind} \\* ARABIC", str(self.counts[kind])),
                p.add_run(f": {title}")]
        for r in runs:
            r.font.size, r.font.bold, r.font.italic, r.font.color.rgb = Pt(10), True, False, NAVY
        if discussion:
            d = self.doc.add_paragraph()
            d.alignment = p.alignment
            d.paragraph_format.space_after = Pt(4 if kind == "Table" else 10)
            d.paragraph_format.keep_with_next = kind == "Table"
            d.paragraph_format.line_spacing = 1.0
            self.inline(d, discussion, size=10)
            if d.runs:
                d.runs[0].text = d.runs[0].text.lstrip()
            for r in d.runs:
                r.font.italic, r.font.color.rgb = True, GREY

    def inline(self, paragraph, children, size=None, bold=False):
        state = {"bold": bold, "italic": False, "code": False}
        for tok in children:
            t = tok.type
            if t in ("strong_open", "strong_close"):
                state["bold"] = t == "strong_open" or bold
            elif t in ("em_open", "em_close"):
                state["italic"] = t == "em_open"
            elif t in ("text", "code_inline", "html_inline"):
                run = paragraph.add_run(tok.content)
                run.bold, run.italic = state["bold"], state["italic"]
                if t == "code_inline":
                    run.font.name = CODE_FONT
                    run.font.size = Pt((size or 12) - 2)
                    run.font.color.rgb = NAVY
                elif size:
                    run.font.size = Pt(size)
            elif t in ("softbreak",):
                paragraph.add_run(" ")
            elif t == "hardbreak":
                paragraph.add_run().add_break()
            elif t == "image":
                self.image(tok.attrs.get("src", ""), tok.content)

    def image(self, src: str, alt: str, caption=None):
        path = (self.base / src).resolve()
        if not path.exists():
            self.doc.add_paragraph(f"[missing image: {src}]")
            return
        for part in split_tall(path):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            pix = pymupdf.Pixmap(str(part))
            # leave room on the page for a heading above and the caption below
            width = min(IMAGE_WIDTH_CM, IMAGE_FIT_HEIGHT_CM * pix.width / pix.height)
            p.add_run().add_picture(str(part), width=Cm(width))
        if caption is None:
            self.uncaptioned.append(f"figure {src}")
            caption = (alt, [])
        self.caption("Figure", *caption)

    def table(self, rows):
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = 9 if ncols >= 5 else 10
        # widths proportional to each column's longest text, so descriptions get the room
        def measure(children):
            """(longest unbreakable word, total length); monospace code counts wider."""
            longest, total = 0, 0
            for c in children:
                if c.type in ("text", "code_inline"):
                    scale = 1.4 if c.type == "code_inline" else 1.0
                    total += len(c.content) * scale
                    longest = max([longest] + [len(w) * scale for w in c.content.split()])
            return longest, total

        weights, floors = [], []
        for j in range(ncols):
            sizes = [measure(r[j]) for r in rows if j < len(r)]
            word = max(w for w, _ in sizes)
            text = min(max(t for _, t in sizes), 60)
            weights.append(max(word + 2, text * 0.6, 5))
            floors.append(min(word * size * 0.024 + 0.4, 4.0))   # cm the longest word needs at this size
        cms = [TEXT_WIDTH_CM * w / sum(weights) for w in weights]
        # widen any column too narrow for its longest word, taking the room from the widest
        for j in range(ncols):
            short = floors[j] - cms[j]
            if short > 0:
                cms[j] += short
                cms[cms.index(max(cms))] -= short
        widths = [Cm(c) for c in cms]
        table.autofit = False
        for i, cells in enumerate(rows):
            row = table.add_row()
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))       # a row never breaks across pages
            if i == 0:
                tr_pr.append(OxmlElement("w:tblHeader"))   # the header repeats on every page
            for j in range(ncols):
                cell = row.cells[j]
                cell.width = widths[j]
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                # short tables stay on one page, so no row is left alone on the next one
                cell.paragraphs[0].paragraph_format.keep_with_next = len(rows) <= 8 and i < len(rows) - 1
                if j < len(cells):
                    self.inline(cell.paragraphs[0], cells[j], size=size, bold=(i == 0))
                if i == 0:
                    shade(cell._tc, HEADER_FILL)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, text: str):
        lines = text.rstrip("\n").split("\n")
        for k, line in enumerate(lines):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = len(lines) <= 25 and k < len(lines) - 1  # short blocks stay whole
            p.paragraph_format.left_indent = Cm(0.3)
            shade(p._p, CODE_FILL)
            r = p.add_run(line if line else " ")
            r.font.name, r.font.size = CODE_FONT, Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def render(self, text: str, shift: int, chapter_title: str | None, chapter: int | None = None):
        tokens = self.md.parse(text)
        i, list_stack, first_h1, pending_table = 0, [], True, None
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "heading_open":
                level = int(tok.tag[1]) + shift
                content = tokens[i + 1].content.replace("`", "")
                if level == 1 and first_h1 and chapter_title:
                    content = chapter_title  # the chapter's own numbered title
                    first_h1 = False
                elif level == 2 and chapter:
                    self.section += 1  # number sections chapter.section, replacing the file's own numbers
                    content = f"{chapter}.{self.section} " + re.sub(r"^\d+\.\s+", "", content)
                self.doc.add_heading(content, level=min(level, 4))
                i += 3
                continue
            if tok.type in ("bullet_list_open", "ordered_list_open"):
                list_stack.append([tok.type == "ordered_list_open", int(tok.attrs.get("start", 1) or 1)])
            elif tok.type in ("bullet_list_close", "ordered_list_close"):
                list_stack.pop()
            elif tok.type == "paragraph_open":
                inline = tokens[i + 1]
                if len(inline.children) == 1 and inline.children[0].type == "image":
                    # the caption is the paragraph right after the image: **Figure: Title.** discussion
                    nxt = tokens[i + 4] if i + 4 < len(tokens) and tokens[i + 3].type == "paragraph_open" else None
                    cap = self.caption_of(nxt.children, "Figure") if nxt else None
                    self.image(inline.children[0].attrs["src"], inline.children[0].content, cap)
                    i += 6 if cap else 3
                    continue
                elif (cap := self.caption_of(inline.children, "Table")) and i + 3 < len(tokens) \
                        and tokens[i + 3].type == "table_open":
                    pending_table = cap  # the caption sits above its table
                elif list_stack and tokens[i - 1].type == "list_item_open" and self.hanging:
                    p = self.doc.add_paragraph()   # a reference entry: hanging indent, no bullet
                    p.paragraph_format.left_indent = Cm(1.27)
                    p.paragraph_format.first_line_indent = Cm(-1.27)
                    p.paragraph_format.line_spacing = 1.15
                    self.inline(p, inline.children)
                elif list_stack and tokens[i - 1].type == "list_item_open":
                    ordered, counter = list_stack[-1]
                    p = self.doc.add_paragraph(style="List Bullet" if not ordered else None)
                    p.paragraph_format.space_after = Pt(3)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    if ordered:
                        p.paragraph_format.left_indent = Cm(0.8)
                        p.paragraph_format.first_line_indent = Cm(-0.5)
                        p.add_run(f"{counter}. ")
                        list_stack[-1][1] += 1
                    self.inline(p, inline.children)
                else:
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    if list_stack:
                        p.paragraph_format.left_indent = Cm(0.8)
                    else:
                        p.paragraph_format.first_line_indent = Cm(1.27)
                    # a sentence introducing a figure or table stays on the page with it
                    nxt = tokens[i + 4] if i + 4 < len(tokens) and tokens[i + 3].type == "paragraph_open" else None
                    p.paragraph_format.keep_with_next = bool(nxt) and (
                        (len(nxt.children) == 1 and nxt.children[0].type == "image")
                        or bool(self.caption_of(nxt.children, "Table")))
                    self.inline(p, inline.children)
                i += 3
                continue
            elif tok.type == "table_open":
                rows, cur = [], None
                i += 1
                while tokens[i].type != "table_close":
                    t = tokens[i]
                    if t.type == "tr_open":
                        cur = []
                    elif t.type == "inline":
                        cur.append(t.children)
                    elif t.type == "tr_close":
                        rows.append(cur)
                    i += 1
                if pending_table is None:
                    first = "".join(c.content for c in rows[0][0]) if rows and rows[0] else ""
                    self.uncaptioned.append(f"table starting '{first}'")
                    pending_table = ("Untitled table", [])
                self.caption("Table", *pending_table)
                pending_table = None
                self.table(rows)
            elif tok.type in ("fence", "code_block"):
                self.code(tok.content)
            i += 1


CAPTION = re.compile(r"^\*\*(Table|Figure): (.+?)\.\*\*", re.M)
REFS: dict[str, str] = {}   # "table-<title slug>" -> "Table 12", filled before rendering


def slug(kind: str, title: str) -> str:
    return kind.lower() + "-" + re.sub(r"[^a-z0-9]+", "-", title.replace("`", "").lower()).strip("-")


def number_captions(names: list[str]) -> None:
    """Number every caption in reading order, so the text can cite "Table 12" before it appears."""
    counts = {"Table": 0, "Figure": 0}
    for name in names:
        for kind, title in CAPTION.findall((DOCS / name).read_text(encoding="utf-8")):
            counts[kind] += 1
            REFS[slug(kind, title)] = f"{kind} {counts[kind]}"


def links_to_text(text: str) -> str:
    """Links to other docs read as the chapter they became; external URLs keep their address."""
    def repl(m):
        label, target = m.group(1), m.group(2)
        if target.startswith("http") and target not in label:
            return f"{label} ({target})"
        if target.startswith("#table-") or target.startswith("#figure-"):
            if target[1:] not in REFS:
                raise SystemExit(f"reference to a missing caption: {target}")
            return REFS[target[1:]]   # [Table](#table-…) prints as "Table 12"
        if target in IN_PRINT:
            return IN_PRINT[target] if label.strip("`") == target else f"{label} ({IN_PRINT[target]})"
        return label
    return re.sub(r"(?<!!)\[([^\]]+)\]\(([^)]+)\)", repl, text)


def sample_queries_appendix(renderer: Renderer) -> None:
    doc = renderer.doc
    doc.add_heading("Appendix A. Sample Analytical Queries", level=1)
    doc.add_paragraph("Each query answers a key business question (section 1.3) and demonstrates an OLAP "
                      "operation. Timings are in the Query Performance Test Report (Chapter 8). Source: "
                      "docs/sample-queries.sql.")
    text = (DOCS / "sample-queries.sql").read_text(encoding="utf-8")
    for part in re.split(r"^-- @query ", text, flags=re.M)[1:]:
        header, _, body = part.partition("\n")
        qid, title, operation = [s.strip() for s in header.split("|")]
        doc.add_heading(f"{qid}. {title}", level=3)
        p = doc.add_paragraph()
        r = p.add_run(f"OLAP operation: {operation}")
        r.font.italic, r.font.color.rgb = True, GREY
        renderer.code(body.strip())


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup(doc)
    footer(doc)
    title_page(doc)
    renderer = Renderer(doc, DOCS)
    numbered = ([(None, c) for c in FRONT_MATTER] + list(enumerate(CHAPTERS, start=1))
                + [(None, c) for c in BACK_MATTER])
    number_captions([name for _, (_, files) in numbered for name, _ in files])
    for number, (title, files) in numbered:
        chapter_title = f"{number}. {title}" if number else title
        renderer.section = 0
        for k, (name, shift) in enumerate(files):
            text = links_to_text((DOCS / name).read_text(encoding="utf-8"))
            renderer.hanging = name == "references.md"
            renderer.render(text, shift, chapter_title if k == 0 else None, number or None)
    sample_queries_appendix(renderer)
    drop_spacers_before_breaks(doc)
    doc.save(DOCX)
    print(f"wrote {DOCX.relative_to(ROOT)}: {renderer.counts['Figure']} figures, {renderer.counts['Table']} tables")
    for item in renderer.uncaptioned:
        print(f"  no caption: {item}")
    finish_with_word()


def drop_spacers_before_breaks(doc: Document) -> None:
    """Remove empty spacer paragraphs right before a chapter heading: when a chapter ends at the
    foot of a page, the spacer alone would spill onto a page of its own, which is then left blank."""
    for p in list(doc.paragraphs):
        nxt = p._p.getnext()
        if (p.text == "" and not p._p.xpath(".//w:drawing | .//w:fldChar") and nxt is not None
                and nxt.tag == qn("w:p") and nxt.xpath("./w:pPr/w:pStyle[@w:val='Heading1']")):
            p._p.getparent().remove(p._p)


def finish_with_word() -> None:
    """Update the contents, lists of figures and tables, and export a PDF with Microsoft Word."""
    import subprocess
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"try {{ $d = $w.Documents.Open('{DOCX}'); $d.Fields.Update() | Out-Null; "
          f"foreach ($t in $d.TablesOfContents) {{ $t.Update() }}; "
          f"foreach ($t in $d.TablesOfFigures) {{ $t.Update() }}; "
          f"$d.Save(); $d.ExportAsFixedFormat('{DOCX.with_suffix('.pdf')}', 17); $d.Close() }} "
          f"finally {{ $w.Quit() }}")
    result = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    if result.returncode != 0:
        print("Word not available: open the .docx, update all fields (Ctrl+A, F9), and save as PDF.")
        print(result.stderr[-500:])
        return
    print(f"wrote {DOCX.with_suffix('.pdf').relative_to(ROOT)} (contents and lists updated)")


if __name__ == "__main__":
    main()
