"""Export the data marts as one formatted Excel workbook, live from the warehouse.

    python ingestion/export_xlsx.py --out driw.xlsx
    python ingestion/export_xlsx.py --event TC2024-KRISTINE --out -      # to stdout (PHP endpoint)

One sheet per mart (filtered to one event when --event is given), an About sheet, and a Data
dictionary sheet. Headers are styled, the top row is frozen with filters on, and numbers carry
formats: pesos, counts, rates as percentages, hours, dates.
"""

import argparse
import io
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import ROOT, connect

sys.path.insert(0, str(ROOT / "docs" / "build"))
from data_dictionary import COLUMNS, TABLES  # noqa: E402  (curated business descriptions)

MARTS = [  # (mart, sheet name, filterable by event)
    ("mart_event_summary", "Event summary", True),
    ("mart_impact_cube", "Impact cube", False),
    ("mart_province_event_impact", "Province impact", True),
    ("mart_evacuation_daily", "Evacuation daily", True),
    ("mart_response_performance", "Response performance", True),
    ("mart_relief_by_org_category", "Relief by organization", True),
]
HEADER_FILL = PatternFill("solid", fgColor="14233A")
NOTICE_FILL = PatternFill("solid", fgColor="FFF7E6")


def number_format(column: str, value) -> str | None:
    if column.endswith("_php"):
        return '#,##0'
    if column.endswith(("_rate", "_ratio")) or column == "occupancy_rate":
        return "0.0%"
    if column.startswith(("median_hours", "p90_hours", "hours_")):
        return "0.00"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "#,##0" if float(value).is_integer() else "#,##0.00"
    return None


def build(event: str | None) -> Workbook:
    wb = Workbook()
    about = wb.active
    about.title = "About"
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    dictionary: list[tuple[str, str, str, str]] = []

    with connect() as conn:
        if event and not conn.execute("select 1 from mart.mart_event_summary where event_code = %s",
                                      (event,)).fetchone():
            raise SystemExit(f"unknown event {event}")
        for mart, sheet_name, by_event in MARTS:
            cur = conn.execute(f"select * from mart.{mart} limit 0")
            ncols = len(cur.description)
            where = " where event_code = %s" if event and by_event else ""
            cur = conn.execute(f"select * from mart.{mart}{where} order by "
                               + ", ".join(str(i) for i in range(1, ncols + 1)), (event,) if where else None)
            columns = [c.name for c in cur.description]
            types = {c.name: conn.adapters.types.get(c.type_code).name if conn.adapters.types.get(c.type_code) else ""
                     for c in cur.description}
            rows = cur.fetchall()
            ws = wb.create_sheet(sheet_name)
            ws.append(columns)
            for r in rows:
                ws.append([float(v) if v.__class__.__name__ == "Decimal" else v for v in r])
            style_sheet(ws, columns, rows)
            for c in columns:
                dictionary.append((sheet_name, c, types[c], COLUMNS.get(c, "")))

    title = "DRIW — Disaster Response Intelligence Warehouse"
    about["A1"] = title
    about["A1"].font = Font(bold=True, size=14)
    lines = [
        ("Generated", generated),
        ("Scope", f"Event {event}" if event else "All events, 2020–2025"),
        ("Source", "Data marts (mart schema) of the DRIW PostgreSQL warehouse"),
        ("Notice", "Event names, dates and PSGC geography are real. All damage, evacuation, relief and "
                   "response figures are synthetic, generated for demonstration."),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        about.cell(row=i, column=1, value=k).font = Font(bold=True)
        about.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    about.cell(row=6, column=2).fill = NOTICE_FILL
    about.cell(row=8, column=1, value="Sheet").font = Font(bold=True)
    about.cell(row=8, column=2, value="Contents").font = Font(bold=True)
    for i, (mart, sheet_name, by_event) in enumerate(MARTS, start=9):
        about.cell(row=i, column=1, value=sheet_name)
        note = TABLES.get(mart, "") + ("" if by_event or not event else " (all events: not filtered)")
        about.cell(row=i, column=2, value=note).alignment = Alignment(wrap_text=True, vertical="top")
    about.cell(row=9 + len(MARTS), column=1, value="Data dictionary")
    about.cell(row=9 + len(MARTS), column=2, value="Every column of every sheet, with its meaning.")
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 100

    dd = wb.create_sheet("Data dictionary")
    dd.append(["Sheet", "Column", "Type", "Description"])
    for row in dictionary:
        dd.append(list(row))
    style_sheet(dd, ["Sheet", "Column", "Type", "Description"], dictionary)
    dd.column_dimensions["D"].width = 90
    return wb


def style_sheet(ws, columns: list[str], rows: list) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for i, name in enumerate(columns, start=1):
        sample = next((r[i - 1] for r in rows if r[i - 1] is not None), None)
        fmt = number_format(name, sample)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
        width = max([len(name)] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, width + 2))


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--event", help="event code, e.g. TC2024-KRISTINE")
    parser.add_argument("--out", required=True, help="output .xlsx path, or - for stdout")
    args = parser.parse_args(argv)
    if args.event and not re.fullmatch(r"TC\d{4}-[A-Z]+", args.event):
        raise SystemExit("event must look like TC2024-KRISTINE")
    wb = build(args.event)
    if args.out == "-":
        buffer = io.BytesIO()
        wb.save(buffer)
        sys.stdout.buffer.write(buffer.getvalue())
    else:
        wb.save(Path(args.out))
        print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
