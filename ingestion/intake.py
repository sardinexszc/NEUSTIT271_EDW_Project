"""Data exchange: Excel templates, column-mapping profiles, validation and submission of reports
from LGUs and national agencies.

    python ingestion/intake.py templates                 # Excel templates + catalog for the web page
    python ingestion/intake.py samples                   # demo files in two LGUs' own layouts
    python ingestion/intake.py validate --type dromic --file report.xlsx [--profile ID] [--name NAME]
    python ingestion/intake.py submit   --type dromic --file report.xlsx [--profile ID] [--name NAME] [--no-load]

Validation runs the warehouse's own data-quality rules, not a copy of them: the uploaded rows are
loaded into a sandbox database (<warehouse>_intake) that holds the same reference data, and the
same dbt staging models run there. The main warehouse is touched only by `submit`, which writes
the file in the standard layout to data/intake/ and reloads the warehouse. Output is JSON (stdout),
which the PHP endpoint web/php/api/intake.php passes to the Data exchange page.
"""

import argparse
import contextlib
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import time
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import INTAKE_DIR, MIGRATIONS_DIR, REFERENCE_DIR, ROOT, connect
from intake_formats import FIXED_LISTS, REPORT_TYPES, ReportType

PROFILES_DIR = ROOT / "ingestion" / "intake_profiles"
TEMPLATES_DIR = ROOT / "web" / "public" / "templates"
CATALOG_PATH = ROOT / "web" / "src" / "data" / "intake_catalog.json"
SAMPLES_DIR = ROOT / "data" / "intake-samples"
RULES_CSV = ROOT / "dbt" / "seeds" / "dq_rules.csv"
DBT_DIR = ROOT / "dbt"

MAX_ROWS = 20_000
CALENDAR = (date(2019, 1, 1), date(2026, 12, 31))    # range of dw.dim_date
TEMPLATE_ROWS = 2_000                                  # rows that carry drop-down lists

# Checks done before the warehouse rules: problems that would otherwise break or distort a load.
SUBMISSION_RULES = {
    "REQUIRED_MISSING": "A required value is blank.",
    "DATE_UNREADABLE": "The date or time cannot be read.",
    "DATE_OUT_OF_RANGE": "The date is outside 2019–2026, the period the warehouse calendar covers.",
    "ID_DUPLICATE_IN_FILE": "The same ID appears on more than one row of this file.",
    "ID_ALREADY_LOADED": "This ID is already in the warehouse from another file.",
    "ALREADY_REPORTED": "This center, day and event is already reported in the warehouse or earlier in this file.",
}


# ----------------------------------------------------------------------------------------- profiles

@dataclass
class Profile:
    id: str
    name: str
    report_type: str
    description: str = ""
    sheet: str | None = None
    header_row: int | None = None                       # 1-based; None = detect
    columns: dict[str, str] = field(default_factory=dict)          # their header -> standard column
    defaults: dict[str, str] = field(default_factory=dict)         # standard column -> value
    value_maps: dict[str, dict[str, str]] = field(default_factory=dict)
    date_format: str | None = None                      # e.g. %d/%m/%Y when days come first


def standard_profile(rt: ReportType) -> Profile:
    return Profile("standard", f"DRIW template ({rt.title})", rt.key,
                   "The downloadable DRIW template, or any file whose headings are the template's.")


def load_profiles() -> list[Profile]:
    profiles = []
    for path in sorted(PROFILES_DIR.glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        profile = Profile(**data)
        if profile.report_type not in REPORT_TYPES:
            raise SystemExit(f"{path.name}: unknown report_type {profile.report_type}")
        profiles.append(profile)
    return profiles


def find_profile(rt: ReportType, profile_id: str | None) -> Profile:
    if not profile_id or profile_id == "standard":
        return standard_profile(rt)
    for p in load_profiles():
        if p.id == profile_id:
            if p.report_type != rt.key:
                raise IntakeError(f"Layout '{p.name}' is for {REPORT_TYPES[p.report_type].title} reports, "
                                  f"not {rt.title} reports.")
            return p
    raise IntakeError(f"Unknown layout '{profile_id}'.")


class IntakeError(Exception):
    """A problem with the whole file (wrong layout, unreadable, missing columns)."""


# ----------------------------------------------------------------------------------------- reading

def norm_header(text) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def header_map(rt: ReportType, profile: Profile) -> dict[str, str]:
    """Normalized header -> standard column: the template's labels and raw names, plus the profile's."""
    mapping = {}
    for c in rt.columns:
        mapping[norm_header(c.label)] = c.name
        mapping[norm_header(c.name)] = c.name
    for theirs, ours in profile.columns.items():
        if ours not in {c.name for c in rt.columns}:
            raise IntakeError(f"Layout '{profile.name}' maps to unknown column '{ours}'.")
        mapping[norm_header(theirs)] = ours
    return mapping


def read_table(path: Path, rt: ReportType, profile: Profile) -> tuple[list[str], list[tuple[int, list]]]:
    """Header cells and (sheet row number, values) of every non-blank row below the header."""
    if path.suffix.lower() == ".xlsx":
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # corrupt or not really an .xlsx
            raise IntakeError(f"The workbook cannot be opened ({type(exc).__name__}).") from exc
        if profile.sheet:
            if profile.sheet not in wb.sheetnames:
                raise IntakeError(f"Sheet '{profile.sheet}' not found; the workbook has: {', '.join(wb.sheetnames)}.")
            ws = wb[profile.sheet]
        else:
            ws = wb["Report"] if "Report" in wb.sheetnames else wb.worksheets[0]
        rows = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
    elif path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        rows = [(i, r) for i, r in enumerate(csv.reader(io.StringIO(text)), start=1)]
    else:
        raise IntakeError("Upload an Excel workbook (.xlsx) or a CSV file.")

    known = header_map(rt, profile)
    if profile.header_row:
        header_index = profile.header_row
    else:  # the first of the top 15 rows where most filled cells are known headings
        header_index = None
        for i, values in rows[:15]:
            filled = [v for v in values if v not in (None, "")]
            if filled and sum(norm_header(v) in known for v in filled) >= max(2, len(filled) // 2):
                header_index = i
                break
        if header_index is None:
            raise IntakeError("No header row with the expected column headings was found. "
                              "Use the DRIW template, or choose the layout that matches your file.")
    by_number = dict(rows)
    if header_index not in by_number:
        raise IntakeError(f"Header row {header_index} is beyond the end of the sheet.")
    header = ["" if v is None else str(v) for v in by_number[header_index]]
    body = [(i, v) for i, v in rows if i > header_index and any(x not in (None, "") for x in v)]
    if len(body) > MAX_ROWS:
        raise IntakeError(f"The file has {len(body):,} rows; the limit is {MAX_ROWS:,} per upload.")
    if not body:
        raise IntakeError("The file has no data rows below the header.")
    return header, body


# ----------------------------------------------------------------------------------------- values

DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d-%b-%Y", "%b %d, %Y", "%B %d, %Y", "%Y/%m/%d"]
TIME_FORMATS = ["%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M %p",
                "%Y-%m-%dT%H:%M", "%Y-%m-%dT%H:%M:%S", "%d-%b-%Y %H:%M"]


def parse_datetime(value, kind: str, date_format: str | None) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial day number
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    text = str(value).strip()
    formats = ([date_format] if date_format else []) + (TIME_FORMATS + DATE_FORMATS if kind == "datetime" else DATE_FORMATS)
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def as_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class Row:
    number: int                       # row number in the uploaded sheet
    values: dict[str, str | None]     # standard column -> text exactly as the raw layer stores it
    issues: list[dict]                # submission-rule issues
    notes: list[str] = field(default_factory=list)   # values translated by the layout profile


def standardize(rt: ReportType, profile: Profile, header: list[str], body: list) -> tuple[list[Row], dict]:
    known = header_map(rt, profile)
    positions: dict[str, int] = {}
    mapped: dict[str, str] = {}
    for i, h in enumerate(header):
        name = known.get(norm_header(h))
        if name and name not in positions:
            positions[name] = i
            mapped[h] = name
    missing = [c for c in rt.columns if c.required and c.name not in positions and c.name not in profile.defaults]
    if missing:
        raise IntakeError("Required columns not found: " + ", ".join(f"“{c.label}”" for c in missing)
                          + f". Headings found: {', '.join(h for h in header if h) or 'none'}.")

    value_maps = {col: {norm_header(k): v for k, v in m.items()} for col, m in profile.value_maps.items()}
    rows = []
    for number, cells in body:
        values: dict[str, str | None] = {}
        issues: list[dict] = []
        notes: list[str] = []
        for c in rt.columns:
            cell = cells[positions[c.name]] if c.name in positions and positions[c.name] < len(cells) else None
            if cell in (None, "") and c.name in profile.defaults:
                cell = profile.defaults[c.name]
            if cell in (None, ""):
                values[c.name] = None
                if c.required:
                    issues.append(issue("REQUIRED_MISSING", c.label))
                continue
            if c.kind in ("date", "datetime"):
                parsed = parse_datetime(cell, c.kind, profile.date_format)
                if parsed is None:
                    values[c.name] = as_text(cell)
                    issues.append(issue("DATE_UNREADABLE", c.label))
                    continue
                if not CALENDAR[0] <= parsed.date() <= CALENDAR[1]:
                    issues.append(issue("DATE_OUT_OF_RANGE", c.label))
                if c.kind == "datetime":
                    values[c.name] = parsed.strftime("%Y-%m-%d %H:%M")
                elif rt.key == "pagasa":
                    values[c.name] = parsed.strftime("%b %d, %Y")      # the bulletin format PAGASA uses
                else:
                    values[c.name] = parsed.strftime("%Y-%m-%d")
                continue
            text = as_text(cell)
            if c.name in value_maps and norm_header(text) in value_maps[c.name]:
                mapped_value = value_maps[c.name][norm_header(text)]
                if mapped_value != text:
                    notes.append(f"{c.label} “{text}” → {mapped_value} (layout)")
                text = mapped_value
            values[c.name] = text
        rows.append(Row(number, values, issues, notes))

    if rt.id_column and rt.key != "ngo":   # NGO rows share their distribution ID by design
        seen: dict[str, int] = {}
        for r in rows:
            key = r.values.get(rt.id_column)
            if key:
                if key in seen:
                    r.issues.append(issue("ID_DUPLICATE_IN_FILE", rt.column(rt.id_column).label,
                                          f"Same ID as row {seen[key]}."))
                else:
                    seen[key] = r.number
    return rows, mapped


def issue(code: str, field_label: str | None = None, detail: str | None = None) -> dict:
    message = SUBMISSION_RULES[code] + (f" {detail}" if detail else "")
    return {"code": code, "field": field_label, "severity": "reject", "message": message,
            "action": "Not submitted", "rule_set": "submission"}


# ----------------------------------------------------------------------------------------- sandbox

def main_db() -> str:
    return os.environ["PGDATABASE"]


def sandbox_db() -> str:
    return f"{main_db()}_intake"


@contextmanager
def intake_lock():
    """One validation or submission at a time: they share the sandbox and the warehouse load."""
    path = Path(tempfile.gettempdir()) / f"driw_{main_db()}_intake.lock"
    deadline = time.time() + 240
    while True:
        try:
            fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            if time.time() - path.stat().st_mtime > 600:   # stale lock from a crashed run
                path.unlink(missing_ok=True)
                continue
            if time.time() > deadline:
                raise IntakeError("Another upload is being processed. Try again in a minute.")
            time.sleep(1)
    try:
        yield
    finally:
        os.close(fd)
        path.unlink(missing_ok=True)


def ensure_sandbox() -> None:
    with connect(autocommit=True, dbname="postgres") as admin:
        if not admin.execute("select 1 from pg_database where datname = %s", (sandbox_db(),)).fetchone():
            admin.execute(f'create database "{sandbox_db()}"')
    # apply every migration the warehouse has (same schemas and raw tables)
    with connect(dbname=sandbox_db()) as conn:
        conn.execute("create schema if not exists meta")
        conn.execute("create table if not exists meta.schema_migration "
                     "(version text primary key, applied_at timestamptz not null default now())")
        applied = {r[0] for r in conn.execute("select version from meta.schema_migration")}
        for path in sorted(MIGRATIONS_DIR.glob("*.sql")):
            if path.stem not in applied:
                conn.execute(path.read_text(encoding="utf-8"))
                conn.execute("insert into meta.schema_migration (version) values (%s)", (path.stem,))
        conn.commit()


REF_TABLES = ["ref.psgc", "ref.organization", "ref.organization_alias", "ref.resource_item", "ref.resource_item_alias"]


def refresh_sandbox(rt: ReportType, rows: list, filename: str) -> int:
    """Copy reference data and known events from the warehouse, then land the upload as one batch."""
    with connect(dbname=main_db()) as src, connect(dbname=sandbox_db()) as dst:
        dst.execute("truncate meta.load_batch cascade")    # empties every raw table too
        dst.execute("truncate " + ", ".join(REF_TABLES) + " cascade")
        for table in REF_TABLES:
            with src.cursor().copy(f"copy {table} to stdout") as out, \
                    dst.cursor().copy(f"copy {table} from stdin") as inp:
                for chunk in out:
                    inp.write(chunk)
        # events the warehouse already knows, so names in the upload resolve the same way
        events = src.execute("select local_name, international_name, category, max_wind_kph, max_signal, "
                             "start_date, end_date, affected_areas, _source_file, _row_number "
                             "from raw.pagasa_event order by _batch_id, _row_number").fetchall()
        batch = dst.execute("insert into meta.load_batch (source_system, source_file) values ('PAGASA', "
                            "'warehouse events') returning batch_id").fetchone()[0]
        with dst.cursor().copy("copy raw.pagasa_event (local_name, international_name, category, max_wind_kph, "
                               "max_signal, start_date, end_date, affected_areas, _source_file, _row_number, "
                               "_batch_id) from stdin") as copy:
            for e in events:
                copy.write_row(list(e) + [batch])
        batch = dst.execute("insert into meta.load_batch (source_system, source_file, origin) values (%s, %s, "
                            "'intake') returning batch_id", (rt.source_system, filename)).fetchone()[0]
        land(dst, rt, rows, batch, filename)
        dst.commit()
        return batch


def land(conn, rt: ReportType, rows: list, batch: int, filename: str) -> None:
    if rt.key == "ngo":
        records = ngo_records(rows)
        with conn.cursor().copy("copy raw.ngo_distribution (record, _batch_id, _source_file, _row_number) "
                                "from stdin") as copy:
            for n, (record, _) in enumerate(records, start=1):
                copy.write_row([json.dumps(record, ensure_ascii=False), batch, filename, n])
        return
    cols = [c.name for c in rt.columns]
    with conn.cursor().copy(f"copy raw.{rt.raw_table} ({', '.join(cols)}, _batch_id, _source_file, "
                            "_row_number) from stdin") as copy:
        for n, r in enumerate(rows, start=1):
            copy.write_row([r.values[c] for c in cols] + [batch, filename, n])


def ngo_records(rows: list) -> list[tuple[dict, list]]:
    """Item rows -> one JSON record per distribution ID (the shape partner APIs send), with the rows behind it."""
    grouped: dict[str, tuple[dict, list]] = {}
    for r in rows:
        v = r.values
        key = v["distribution_id"]
        if key not in grouped:
            location = {"province": v["province"]}
            if v["municipality"]:
                location["municipality"] = v["municipality"]
            record = {
                "distribution_id": key, "organization": v["organization"], "event_ref": v["event_ref"],
                "distributed_on": v["distributed_on"], "location": location,
                "beneficiaries": {"families": int(float(v["beneficiary_families"]))
                                  if (v["beneficiary_families"] or "").replace(".", "", 1).isdigit() else None},
                "items": [],
            }
            grouped[key] = (record, [])
        quantity = v["quantity"]
        try:
            quantity = float(str(quantity).replace(",", ""))
            quantity = int(quantity) if quantity.is_integer() else quantity
        except ValueError:
            pass
        grouped[key][0]["items"].append({"item_code": v["item_code"], "unit": v["unit"], "quantity": quantity})
        grouped[key][1].append(r)
    return list(grouped.values())


def run_staging(rt: ReportType) -> None:
    """The warehouse's own staging models (and only what they depend on), in the sandbox."""
    from dbt.cli.main import dbtRunner   # imported here: slow to load, unused by the other commands

    work = Path(tempfile.gettempdir()) / f"driw_{main_db()}_intake_dbt"
    previous = os.environ["PGDATABASE"]
    os.environ["PGDATABASE"] = sandbox_db()
    try:
        # a second attempt re-parses the project from scratch, in case dbt's cached parse is stale
        for extra in ([], ["--no-partial-parse"]):
            # dbt logs to stdout, which must stay clean JSON for the PHP endpoint
            with contextlib.redirect_stdout(sys.stderr):
                result = dbtRunner().invoke([
                    "run", "--select", f"+{rt.staging_model}", "--project-dir", str(DBT_DIR), "--profiles-dir",
                    str(DBT_DIR), "--target-path", str(work / "target"), "--log-path", str(work / "logs"),
                    "--quiet", *extra,
                ])
            if result.success:
                return
            reason = result.exception or [r.message for r in (result.result or []) if r.status == "error"]
            print(f"intake: dbt run failed ({'retrying' if not extra else 'giving up'}): {reason}", file=sys.stderr)
    finally:
        os.environ["PGDATABASE"] = previous
    raise IntakeError("The data-quality check could not run on this file. Please contact the DRIW administrator.")


# what each staging model reveals about how a row was understood
RESULT_SQL = {
    "dromic": """select s._row_number, null::int, s.dq_status, s.dq_issues,
                        s.raw_municipality, p.name, s.location_match_method, s.event_key,
                        s.raw_ec_code, s.ec_code, s.report_date, null, null
                 from staging.stg_dromic_evacuation s left join ref.psgc p using (psgc_code)
                 where s._batch_id = %s""",
    "ndrrmc": """select s._row_number, null::int, s.dq_status, s.dq_issues,
                        s.raw_city_muni, p.name, s.location_match_method, s.event_key,
                        null, null, null, null, null
                 from staging.stg_ndrrmc_damage s left join ref.psgc p using (psgc_code)
                 where s._batch_id = %s""",
    "lgu": """select s._row_number, null::int, s.dq_status, s.dq_issues,
                     s.raw_lgu_name, p.name, s.location_match_method, s.event_key,
                     s.raw_request_type, s.request_type, null, s.responding_unit, s.org_code
              from staging.stg_lgu_response s left join ref.psgc p using (psgc_code)
              where s._batch_id = %s""",
    "pagasa": """select s._row_number, null::int, s.dq_status, s.dq_issues,
                        null, null, null, s.event_key, null, null, null, null, null
                 from staging.stg_pagasa_event s where s._batch_id = %s""",
    "ngo": """select s._row_number, s.line_no, s.dq_status, s.dq_issues,
                     s.raw_municipality, p.name, s.location_match_method, s.event_key,
                     s.source_item_code || ' ' || s.source_unit, s.quantity_standard || ' ' || s.standard_unit,
                     s.to_standard_factor, s.raw_organization, o.org_name
              from staging.stg_ngo_distribution_line s
              left join ref.psgc p using (psgc_code)
              left join ref.organization o using (org_code)
              where s._batch_id = %s""",
}


def warehouse_rules() -> dict[tuple[str, str], dict]:
    with open(RULES_CSV, newline="", encoding="utf-8") as f:
        return {(r["source_system"], r["rule_code"]): r for r in csv.DictReader(f)}


def explain(rt: ReportType, raw_issues: list[str], rules: dict) -> list[dict]:
    out = []
    for code in raw_issues or []:
        rule_code, _, field_name = code.partition(":")
        rule = rules.get((rt.rules_source, rule_code), {})
        label = next((c.label for c in rt.columns if c.name == field_name), field_name or None)
        out.append({"code": rule_code, "field": label, "severity": rule.get("severity", "info"),
                    "message": rule.get("description", rule_code), "action": rule.get("action", ""),
                    "rule_set": "warehouse"})
    return out


# ----------------------------------------------------------------------------------------- validate

STATUS_ORDER = ["accepted", "corrected", "warning", "duplicate", "rejected"]
NOT_CORRECTIONS = {"REQUEST_OPEN"}   # informational: an open request is accepted as it is


def validate(rt: ReportType, profile: Profile, path: Path, filename: str) -> dict:
    header, body = read_table(path, rt, profile)
    rows, mapped = standardize(rt, profile, header, body)
    with connect(dbname=main_db()) as conn:
        check_ids(conn, rt, rows, stored_name(rt, profile, filename))

    ensure_sandbox()
    batch = refresh_sandbox(rt, rows, filename)
    run_staging(rt)
    rules = warehouse_rules()
    with connect(dbname=sandbox_db()) as conn:
        results = conn.execute(RESULT_SQL[rt.key], (batch,)).fetchall()

    # NGO staging rows are item lines of records; map them back to the sheet rows
    if rt.key == "ngo":
        sheet_row = {}
        for n, (_, members) in enumerate(ngo_records(rows), start=1):
            for line, r in enumerate(members, start=1):
                sheet_row[(n, line)] = r
    else:
        sheet_row = {(n, None): r for n, r in enumerate(rows, start=1)}

    report_rows = []
    for (row_n, line, status, raw_issues, raw_place, official, method, event_key, a, b, c, d, e) in results:
        r = sheet_row[(row_n, line)]
        issues = r.issues + explain(rt, raw_issues, rules)
        resolved = list(r.notes)
        if method in ("normalized", "fuzzy") and official:
            resolved.append(f"City/Municipality “{raw_place.strip()}” → {official}")
        if rt.key == "dromic" and a is None and b:
            resolved.append(f"Evacuation center code recovered: {b}")
        if rt.key == "lgu":
            if a and b and a.upper() != b:
                resolved.append(f"Request type “{a}” → {b}")
        if rt.key == "ngo":
            if e and d and d != e:
                resolved.append(f"Organization “{d}” → {e}")
            if c is not None and float(c) != 1 and b:
                resolved.append(f"Quantity {a} → {b}")
        if event_key and status != "rejected":
            resolved.append(f"Event → {event_key}")
        r.values["_event_key"] = event_key
        r.values["_ec_code"] = b if rt.key == "dromic" else None
        report_rows.append({"row": r.number, "dq_status": status, "issues": issues, "resolved": resolved, "_row": r})

    if rt.key == "dromic":
        check_already_reported(rt, report_rows, stored_name(rt, profile, filename))

    for rr in report_rows:
        if any(i["rule_set"] == "submission" for i in rr["issues"]):
            rr["status"] = "rejected"
        elif rr["dq_status"] == "duplicate":
            rr["status"] = "duplicate"
        elif rr["dq_status"] == "rejected":
            rr["status"] = "rejected"
        elif rr["dq_status"] == "warning":
            rr["status"] = "warning"
        # valid, but something was fixed on the way in (a name matched, a code mapped, a unit converted)
        elif (any(i["code"] not in NOT_CORRECTIONS for i in rr["issues"])
              or any(not x.startswith("Event →") for x in rr["resolved"])):
            rr["status"] = "corrected"
        else:
            rr["status"] = "accepted"

    report_rows.sort(key=lambda rr: (rr["row"]))
    summary = {s: sum(rr["status"] == s for rr in report_rows) for s in STATUS_ORDER}
    summary["rows"] = len(report_rows)
    return {
        "ok": True,
        "report_type": rt.key, "report_title": rt.title,
        "profile": {"id": profile.id, "name": profile.name},
        "file": filename,
        "columns_mapped": mapped,
        "summary": summary,
        "rows": [{k: v for k, v in rr.items() if not k.startswith("_") and k != "dq_status"} for rr in report_rows],
        "_report_rows": report_rows,
    }


def check_ids(conn, rt: ReportType, rows: list, target: str) -> None:
    """IDs must be new to the warehouse (a resubmitted file may reuse its own IDs: it replaces itself)."""
    if not rt.id_column:
        return
    ids = sorted({r.values[rt.id_column] for r in rows if r.values.get(rt.id_column)})
    if not ids:
        return
    if rt.key == "ngo":
        sql = ("select record ->> 'distribution_id', _source_file from raw.ngo_distribution "
               "where record ->> 'distribution_id' = any(%s)")
    else:
        sql = f"select {rt.id_column}, _source_file from raw.{rt.raw_table} where {rt.id_column} = any(%s)"
    taken = {key: f for key, f in conn.execute(sql, (ids,)).fetchall() if f != target}
    for r in rows:
        key = r.values.get(rt.id_column)
        if key in taken:
            r.issues.append(issue("ID_ALREADY_LOADED", rt.column(rt.id_column).label, f"Loaded from {taken[key]}."))


def check_already_reported(rt: ReportType, report_rows: list, target: str) -> None:
    """Evacuation snapshots: one row per center x day x event across the whole warehouse."""
    keys = [(rr["_row"].values["_ec_code"], rr["_row"].values["report_date"], rr["_row"].values["_event_key"])
            for rr in report_rows]
    wanted = [k for k in keys if all(k)]
    taken = set()
    if wanted:
        with connect(dbname=main_db()) as conn:
            taken = set(conn.execute(
                "select ec_code, report_date::text, event_key from staging.stg_dromic_evacuation "
                "where dq_status in ('valid', 'warning') and _source_file <> %s "
                "and (ec_code, report_date::text, event_key) in (select * from unnest(%s::text[], %s::text[], %s::text[]))",
                (target, [k[0] for k in wanted], [k[1] for k in wanted], [k[2] for k in wanted])).fetchall())
    seen = set()
    for rr, k in zip(report_rows, keys):
        if not all(k) or rr["dq_status"] not in ("valid", "warning"):
            continue
        if k in taken or k in seen:
            rr["issues"].append(issue("ALREADY_REPORTED", "Evacuation center code"))
        seen.add(k)


# ----------------------------------------------------------------------------------------- submit

def target_name(profile: Profile, filename: str) -> str:
    stem = re.sub(r"[^a-z0-9]+", "-", Path(filename).stem.lower()).strip("-")[:60] or "upload"
    return f"intake_{profile.id}_{stem}"


def stored_name(rt: ReportType, profile: Profile, filename: str) -> str:
    """Name of the file in data/intake/ (and of its batches' _source_file): a resubmission replaces it."""
    return f"{target_name(profile, filename)}.{rt.output}"


def submit(rt: ReportType, profile: Profile, path: Path, filename: str, load: bool) -> dict:
    report = validate(rt, profile, path, filename)
    keep = [rr["_row"] for rr in report["_report_rows"] if not any(i["rule_set"] == "submission" for i in rr["issues"])]
    excluded = len(report["_report_rows"]) - len(keep)
    folder = INTAKE_DIR / rt.key
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / stored_name(rt, profile, filename)
    for old in folder.glob(f"{target_name(profile, filename)}.*"):
        old.unlink()   # a resubmitted file replaces its earlier version
    if keep:
        write_standard(rt, keep, out)
    report["submitted"] = {"file": f"{rt.key}/{out.name}" if keep else None,   # relative to the intake folder
                           "rows_written": len(keep), "rows_excluded": excluded}
    if load:
        report["submitted"]["load"] = load_warehouse()
    return report


def write_standard(rt: ReportType, rows: list, out: Path) -> None:
    """The standard layout data/intake/ is loaded from (the same readers as the synthetic extracts)."""
    cols = [c.name for c in rt.columns]
    if rt.output == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.append(cols)
        for r in rows:
            ws.append([r.values[c] for c in cols])
        wb.save(out)
    elif rt.output == "csv":
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(cols)
            for r in rows:
                w.writerow(["" if r.values[c] is None else r.values[c] for c in cols])
    else:
        out.write_text(json.dumps([rec for rec, _ in ngo_records(rows)], ensure_ascii=False, indent=1),
                       encoding="utf-8")


def load_warehouse() -> dict:
    """Reload raw (synthetic + intake files) and rebuild the warehouse with all its tests."""
    started = time.time()
    result = subprocess.run([sys.executable, str(ROOT / "ingestion" / "pipeline.py"), "raw", "dbt"],
                            cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace")
    done = re.findall(r"Done\. PASS=(\d+) WARN=(\d+) ERROR=(\d+)", result.stdout)
    return {"ok": result.returncode == 0, "seconds": round(time.time() - started),
            "message": (f"Warehouse rebuilt: {done[-1][0]} checks passed." if result.returncode == 0 and done
                        else "The warehouse load failed; the file is kept in data/intake/ for the administrator.")}


# ----------------------------------------------------------------------------------------- templates

def reference_lists() -> dict[str, list[str]]:
    with connect(dbname=main_db()) as conn:
        q = lambda sql: [r[0] for r in conn.execute(sql).fetchall()]  # noqa: E731
        lists = {
            "regions": q("select distinct name from ref.psgc where geographic_level = 'Reg' order by 1"),
            "provinces": q("select distinct province_name from ref.psgc where geographic_level in ('City', 'Mun') order by 1"),
            "cities": q("select distinct name from ref.psgc where geographic_level in ('City', 'Mun') order by 1"),
            "events": q("select distinct initcap(local_name) from raw.pagasa_event order by 1"),
            "incidents": q("select distinct trim(category) || ' ' || initcap(local_name) || ' (' || "
                           "initcap(international_name) || ')' from raw.pagasa_event order by 1"),
            "organizations": q("select org_name from ref.organization order by 1"),
        }
    with open(REFERENCE_DIR / "organizations.csv", newline="", encoding="utf-8") as f:
        lists["units"] = sorted({a.strip() for r in csv.DictReader(f) for a in r["source_aliases"].split(";")})
    with open(REFERENCE_DIR / "resource_items.csv", newline="", encoding="utf-8") as f:
        items = list(csv.DictReader(f))
    lists["item_codes"] = [i["item_code"] for i in items]
    lists["units_of_measure"] = sorted({i["standard_unit"] for i in items}
                                       | {a.split(":")[1] for i in items for a in i["source_aliases"].split(";")})
    lists.update(FIXED_LISTS)
    return lists


HEADER_FILL = PatternFill("solid", fgColor="14233A")
REQUIRED_FILL = PatternFill("solid", fgColor="B9471A")


def build_template(rt: ReportType, lists: dict[str, list[str]], out: Path) -> None:
    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    report = wb.create_sheet("Report")
    lst = wb.create_sheet("Lists")

    info["A1"] = f"DRIW template: {rt.title}"
    info["A1"].font = Font(bold=True, size=14)
    info["A2"] = f"For: {rt.issuer}. {rt.description}"
    info["A3"] = ("Fill in the Report sheet, one row per record, then upload the file on the DRIW Data exchange "
                  "page. Orange headings are required. Drop-down lists offer the official names; other spellings "
                  "are still accepted and matched automatically. Do not rename or reorder the headings.")
    info["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    info.row_dimensions[3].height = 48
    info.append([])
    info.append(["Column", "Required", "What to enter", "Example"])
    for cell in info[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for c in rt.columns:
        info.append([c.label, "Yes" if c.required else "No", c.help, c.example])
    for col, width in zip("ABCD", (28, 10, 70, 34)):
        info.column_dimensions[col].width = width
    for row in info.iter_rows(min_row=6):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    for i, c in enumerate(rt.columns, start=1):
        cell = report.cell(row=1, column=i, value=c.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if c.required else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = get_column_letter(i)
        report.column_dimensions[letter].width = max(14, min(36, len(c.label) + 6))
        if c.kind in ("date", "datetime"):
            for r in range(2, TEMPLATE_ROWS + 2):
                report.cell(row=r, column=i).number_format = "yyyy-mm-dd" if c.kind == "date" else "yyyy-mm-dd hh:mm"
    report.freeze_panes = "A2"
    report.row_dimensions[1].height = 32

    # Lists sheet: one column per list used by this report type; drop-downs point at it
    used = [c for c in rt.columns if c.choices]
    for i, name in enumerate(dict.fromkeys(c.choices for c in used), start=1):
        values = lists[name]
        lst.cell(row=1, column=i, value=name)
        for r, v in enumerate(values, start=2):
            lst.cell(row=r, column=i, value=v)
        letter = get_column_letter(i)
        ref = f"Lists!${letter}$2:${letter}${len(values) + 1}"
        for c in (c for c in used if c.choices == name):
            target = get_column_letter([x.name for x in rt.columns].index(c.name) + 1)
            dv = DataValidation(type="list", formula1=ref, allow_blank=True, errorStyle="warning",
                                error="Not on the official list. DRIW will try to match it; check the spelling.",
                                errorTitle="DRIW")
            dv.add(f"{target}2:{target}{TEMPLATE_ROWS + 1}")
            report.add_data_validation(dv)
    lst.sheet_state = "hidden"
    wb.active = 1
    wb.save(out)


def write_templates() -> None:
    lists = reference_lists()
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    for rt in REPORT_TYPES.values():
        out = TEMPLATES_DIR / f"driw-{rt.key}-template.xlsx"
        build_template(rt, lists, out)
        print(f"wrote {out.relative_to(ROOT)}")
    catalog = {
        "report_types": [{
            "key": rt.key, "title": rt.title, "issuer": rt.issuer, "description": rt.description,
            "template": f"templates/driw-{rt.key}-template.xlsx",
            "columns": [{"label": c.label, "required": c.required} for c in rt.columns],
        } for rt in REPORT_TYPES.values()],
        "profiles": [{"id": p.id, "name": p.name, "report_type": p.report_type, "description": p.description}
                     for p in load_profiles()],
        "submission_rules": SUBMISSION_RULES,
    }
    CATALOG_PATH.write_text(json.dumps(catalog, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {CATALOG_PATH.relative_to(ROOT)}")


# ----------------------------------------------------------------------------------------- samples

def write_samples() -> None:
    """Two demo reports in the LGUs' own layouts (see ingestion/intake_profiles/), with realistic errors."""
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "EC Report"
    ws.append(["PROVINCE OF ALBAY - PROVINCIAL DISASTER RISK REDUCTION AND MANAGEMENT OFFICE"])
    ws.append(["Evacuation Center Status Report - Severe Tropical Storm Kristine (Trami)"])
    ws.append([])
    ws.append(["Date of Report", "Typhoon", "Town/City", "Barangay", "Evacuation Center", "EC ID",
               "Type of Facility", "Capacity (fam)", "No. of Families", "No. of Persons"])
    rows = [
        ("10/22/2024", "Kristine", "Legaspi City", "Bitano", "Bitano Elementary School", "EC-ALB-001", "School", 120, 96, 431),
        ("10/23/2024", "Kristine", "Legaspi City", "Bitano", "Bitano Elementary School", "EC-ALB-001", "School", 120, 71, 318),
        ("10/22/2024", "Kristine", "DARAGA", "Bagumbayan", "Daraga National High School", "EC-ALB-002", "School", 200, 150, 642),
        ("10/23/2024", "Kristine", "DARAGA", "Bagumbayan", "Daraga National High School", None, "School", 200, 118, 505),
        ("10/22/2024", "Kristine", "Guinobatan", "Poblacion", "Guinobatan Gymnasium", "EC-ALB-003", "Gymnasium", 80, 97, 415),
        ("10/22/2024", "Kristine", "Sto. Domingo", "San Isidro", "San Isidro Covered Court", "EC-ALB-004", "Covered Court", 60, 44, 190),
        ("10/22/2024", "Kristine", "Ligao City", "Tuburan", "Tuburan Barangay Hall", "EC-ALB-005", "Barangay Hall", 40, 31, 133),
        ("10/22/2024", "Kristine", "Malniao", "Poblacion", "Malinao Central School", "EC-ALB-006", "School", 150, 88, 377),
        ("10/22/2024", "Kristine", "Camalig", "Ilawod", "Ilawod Elementary School", "EC-ALB-007", "School", 90, 64, 61),
        ("10/32/2024", "Kristine", "Tabaco City", "Bombon", "Bombon Parish Church", "EC-ALB-008", "Church", 70, 52, 228),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(SAMPLES_DIR / "albay-pdrrmo-kristine-evacuation.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Incidents"
    ws.append(["Ticket #", "Event", "City", "Type of request", "Priority", "Received", "Dispatched",
               "On scene", "Closed", "Team", "No. of personnel"])
    rows = [
        ("QC-2024-0001", "Carina", "Quezon City", "Rescue", "Critical", "2024-07-24 06:10", "2024-07-24 06:25", "2024-07-24 07:05", "2024-07-24 09:40", "QC Rescue", 12),
        ("QC-2024-0002", "Carina", "QUEZON CITY", "Medical", "Critical", "2024-07-24 06:42", "2024-07-24 06:55", "2024-07-24 07:30", "2024-07-24 08:15", "DOH-HEMS", 4),
        ("QC-2024-0003", "Carina", "Quezon Ctiy", "Evacuation", "Urgent", "2024-07-24 07:20", "2024-07-24 08:05", "2024-07-24 08:50", "2024-07-24 13:00", "QC Rescue", 9),
        ("QC-2024-0004", "Carina", "Quezon City", "Road clearing", "Routine", "2024-07-24 08:00", "2024-07-24 13:30", "2024-07-24 14:20", "2024-07-25 02:00", "DPWH", 6),
        ("QC-2024-0005", "Carina", "Quezon City", "Relief", "Routine", "2024-07-24 09:15", "2024-07-24 18:40", "2024-07-24 19:30", None, "DSWD", 5),
        ("QC-2024-0006", "Carina", "Quezon City", "Rescue", "Critical", "2024-07-24 10:05", "2024-07-24 10:20", "2024-07-24 10:12", "2024-07-24 12:00", "BFP", 8),
        ("QC-2024-0007", "Carina", "Quezon City", "Others", "Routine", "2024-07-24 11:00", "2024-07-24 12:30", "2024-07-24 13:10", "2024-07-24 14:00", "QC Rescue", 3),
        ("QC-2024-0008", "Carina", "Quezon City", "Rescue", "Urgent", "2024-07-24 12:40", "2024-07-24 13:05", "2024-07-24 13:45", "2024-07-24 16:20", "Barangay Tanod", 6),
        ("QC-2024-0008", "Carina", "Quezon City", "Medical", "Urgent", "2024-07-24 13:10", "2024-07-24 13:30", "2024-07-24 14:00", "2024-07-24 15:10", "DOH-HEMS", 3),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(SAMPLES_DIR / "quezon-city-cdrrmo-carina-incidents.xlsx")
    print(f"wrote 2 samples to {SAMPLES_DIR.relative_to(ROOT)}")


# ----------------------------------------------------------------------------------------- CLI

def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="DRIW data exchange")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("templates")
    sub.add_parser("samples")
    for name in ("validate", "submit"):
        p = sub.add_parser(name)
        p.add_argument("--type", required=True, choices=sorted(REPORT_TYPES))
        p.add_argument("--file", required=True, type=Path)
        p.add_argument("--profile", default="standard")
        p.add_argument("--name", help="original file name (defaults to the file's)")
        if name == "submit":
            p.add_argument("--no-load", action="store_true", help="write the file but do not reload the warehouse")
    args = parser.parse_args(argv)

    if args.command == "templates":
        write_templates()
        return 0
    if args.command == "samples":
        write_samples()
        return 0

    rt = REPORT_TYPES[args.type]
    filename = args.name or args.file.name
    try:
        profile = find_profile(rt, args.profile)
        with intake_lock():
            if args.command == "validate":
                report = validate(rt, profile, args.file, filename)
            else:
                report = submit(rt, profile, args.file, filename, load=not args.no_load)
        report.pop("_report_rows", None)
        report["action"] = args.command
    except IntakeError as exc:
        report = {"ok": False, "action": args.command, "report_type": rt.key, "file": filename,
                  "file_errors": [str(exc)]}
    # bytes, not text: a Windows pipe would otherwise encode as cp1252 and fail on "→"
    sys.stdout.buffer.write(json.dumps(report, ensure_ascii=False, default=str).encode("utf-8"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
