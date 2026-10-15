"""End-to-end validation: the pipeline builds a complete, consistent warehouse from scratch."""

import hashlib
import json
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from conftest import ROOT, scalar

sys.path.insert(0, str(ROOT / "ingestion"))

FACTS = ["fact_damage_assessment", "fact_evacuation_snapshot", "fact_resource_distribution", "fact_response_action"]
MARTS = ["mart_event_summary", "mart_impact_cube", "mart_province_event_impact",
         "mart_evacuation_daily", "mart_response_performance", "mart_relief_by_org_category"]


def test_pipeline_ran_every_step(warehouse):
    if not warehouse["pipeline_output"]:
        pytest.skip("testing an existing database (DRIW_TEST_DB)")
    for step in ["migrate", "reference", "psgc", "generate", "raw", "dbt", "export"]:
        assert f"=== {step} ===" in warehouse["pipeline_output"]


def test_every_planted_defect_is_detected_exactly(db):
    rows = db.execute("select planted_defect, planted_count, detected_count, reconciliation_status "
                      "from dq.dq_reconciliation").fetchall()
    assert len(rows) == 19
    mismatches = [r for r in rows if r[3] != "match"]
    assert not mismatches, mismatches


@pytest.mark.parametrize("schema,table", [("raw", t) for t in
                         ["pagasa_event", "ndrrmc_damage_report", "dromic_evacuation", "ngo_distribution", "lgu_response_log"]]
                         + [("dw", f) for f in FACTS] + [("mart", m) for m in MARTS])
def test_no_layer_is_empty(db, schema, table):
    assert scalar(db, f"select count(*) from {schema}.{table}") > 0


@pytest.mark.parametrize("raw_table,staging_table", [
    ("pagasa_event", "stg_pagasa_event"),
    ("ndrrmc_damage_report", "stg_ndrrmc_damage"),
    ("dromic_evacuation", "stg_dromic_evacuation"),
    ("lgu_response_log", "stg_lgu_response"),
])
def test_staging_keeps_every_raw_row(db, raw_table, staging_table):
    """Data quality flags rows; it never silently drops them."""
    assert scalar(db, f"select count(*) from raw.{raw_table}") == scalar(db, f"select count(*) from staging.{staging_table}")


def test_staging_keeps_every_ngo_item_line(db):
    raw_lines = scalar(db, "select sum(jsonb_array_length(record -> 'items')) from raw.ngo_distribution")
    assert raw_lines == scalar(db, "select count(*) from staging.stg_ngo_distribution_line")


def test_rows_are_accounted_for(db):
    """Every staged row is valid, warning, duplicate or rejected, and every rejected row is quarantined."""
    for model, source in [("stg_ndrrmc_damage", "ndrrmc"), ("stg_dromic_evacuation", "dromic"),
                          ("stg_ngo_distribution_line", "ngo"), ("stg_lgu_response", "lgu")]:
        rejected = scalar(db, f"select count(*) from staging.{model} where dq_status = 'rejected'")
        quarantined = scalar(db, "select count(*) from dq.dq_quarantine where source_system = %s", (source,))
        assert rejected == quarantined, model
        assert scalar(db, f"select count(*) from staging.{model} where dq_status is null") == 0


def test_damage_totals_agree_across_layers(db):
    fact = scalar(db, "select sum(cost_php) from dw.fact_damage_assessment")
    assert fact == scalar(db, "select sum(damage_cost_php) from mart.mart_event_summary")
    assert fact == scalar(db, "select damage_cost_php from mart.mart_impact_cube where grouping_level = 'grand_total'")


def test_exported_json_matches_marts(warehouse, db):
    for mart in MARTS:
        path = warehouse["export_dir"] / f"{mart}.json"
        assert path.exists(), path
        rows = json.loads(path.read_text(encoding="utf-8"))
        assert len(rows) == scalar(db, f"select count(*) from mart.{mart}"), mart


def _content_digest(folder):
    """Hash of every generated file's content (xlsx compared by cell values, not zip metadata)."""
    h = hashlib.sha256()
    for path in sorted(folder.rglob("*")):
        if path.is_dir():
            continue
        h.update(path.relative_to(folder).as_posix().encode())
        if path.suffix == ".xlsx":
            for row in load_workbook(path, read_only=True).active.iter_rows(values_only=True):
                h.update(repr(row).encode())
        else:
            h.update(path.read_bytes())
    return h.hexdigest()


def test_generator_is_deterministic(tmp_path):
    import generate_sources

    digests = []
    for run in ("a", "b"):
        generate_sources.SOURCES_DIR = tmp_path / run
        generate_sources.generate(seed=271)
        digests.append(_content_digest(tmp_path / run))
    assert digests[0] == digests[1]


def test_rebuild_keeps_security_and_counts(warehouse, db):
    """A second dbt build must recreate the same facts AND re-apply grants and row-level security."""
    if "env" not in warehouse:
        pytest.skip("only run against the throwaway e2e database")
    before = {f: scalar(db, f"select count(*) from dw.{f}") for f in FACTS}
    db.rollback()
    result = subprocess.run([sys.executable, str(ROOT / "ingestion" / "pipeline.py"), "dbt"],
                            cwd=ROOT, env=warehouse["env"], capture_output=True, text=True,
                            encoding="utf-8", errors="replace")
    assert result.returncode == 0, result.stdout[-3000:]
    for f in FACTS:
        assert scalar(db, f"select count(*) from dw.{f}") == before[f]
        assert scalar(db, "select relrowsecurity from pg_class where oid = %s::regclass", (f"dw.{f}",))
        assert scalar(db, "select count(*) from pg_policies where schemaname = 'dw' and tablename = %s", (f,)) == 2
        assert scalar(db, "select has_table_privilege('driw_analyst', %s, 'SELECT')", (f"dw.{f}",))
