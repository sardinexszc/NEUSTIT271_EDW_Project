"""Data exchange: templates, layout profiles, the per-row validation report, submission into the
warehouse, and the Excel export.

Validation and submission run against the throwaway e2e warehouse (and its own sandbox); the
submitted file goes to the e2e intake folder, never the working copy's data/intake/.
"""

import json
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from conftest import ROOT, connect, scalar

sys.path.insert(0, str(ROOT / "ingestion"))
from intake_formats import REPORT_TYPES  # noqa: E402

SAMPLES = ROOT / "data" / "intake-samples"
ALBAY = SAMPLES / "albay-pdrrmo-kristine-evacuation.xlsx"
QUEZON_CITY = SAMPLES / "quezon-city-cdrrmo-carina-incidents.xlsx"


def run_intake(warehouse, *args) -> dict:
    if "env" not in warehouse:
        pytest.skip("only run against the throwaway e2e database")
    result = subprocess.run([sys.executable, str(ROOT / "ingestion" / "intake.py"), *args], cwd=ROOT,
                            env=warehouse["env"], capture_output=True, timeout=600)
    assert result.returncode == 0, result.stderr.decode("utf-8", "replace")[-3000:]
    return json.loads(result.stdout.decode("utf-8"))


def by_row(report: dict) -> dict[int, dict]:
    return {r["row"]: r for r in report["rows"]}


def codes(row: dict) -> set[str]:
    return {i["code"] for i in row["issues"]}


# ---------------------------------------------------------------------------- templates

@pytest.mark.parametrize("key", sorted(REPORT_TYPES))
def test_template_has_instructions_headings_and_dropdowns(key):
    rt = REPORT_TYPES[key]
    wb = load_workbook(ROOT / "web" / "public" / "templates" / f"driw-{key}-template.xlsx")
    assert wb.sheetnames == ["Instructions", "Report", "Lists"]
    headings = [c.value for c in wb["Report"][1]]
    assert headings == [c.label for c in rt.columns]
    assert wb["Lists"].sheet_state == "hidden"
    if any(c.choices for c in rt.columns):
        assert wb["Report"].data_validations.dataValidation, "drop-down lists missing"


def test_catalog_lists_every_report_type_and_layout():
    catalog = json.loads((ROOT / "web" / "src" / "data" / "intake_catalog.json").read_text(encoding="utf-8"))
    assert [t["key"] for t in catalog["report_types"]] == list(REPORT_TYPES)
    profiles = {p.stem for p in (ROOT / "ingestion" / "intake_profiles").glob("*.json")}
    assert {p["id"] for p in catalog["profiles"]} == profiles


# ---------------------------------------------------------------------------- validation report

def test_albay_layout_report_explains_every_row(warehouse):
    report = run_intake(warehouse, "validate", "--type", "dromic", "--profile", "albay-pdrrmo-evacuation",
                        "--file", str(ALBAY))
    assert report["ok"]
    rows = by_row(report)
    assert report["summary"] == {"accepted": 0, "corrected": 5, "warning": 2, "duplicate": 0, "rejected": 3, "rows": 10}
    assert rows[5]["status"] == "corrected" and any("City of Legazpi" in x for x in rows[5]["resolved"])
    assert any("Santo Domingo" in x for x in rows[10]["resolved"])               # Sto. Domingo
    assert rows[8]["status"] == "warning" and "EC_CODE_IMPUTED" in codes(rows[8])
    assert rows[9]["status"] == "warning" and "EC_OVER_CAPACITY" in codes(rows[9])
    assert rows[12]["status"] == "rejected" and "LOC_UNMATCHED" in codes(rows[12])     # Malniao: a tie
    assert rows[13]["status"] == "rejected" and "PERSONS_LT_FAMILIES" in codes(rows[13])
    assert rows[14]["status"] == "rejected" and "DATE_UNREADABLE" in codes(rows[14])   # 10/32/2024


def test_validation_does_not_touch_the_warehouse(warehouse, db):
    before = scalar(db, "select count(*) from raw.dromic_evacuation")
    run_intake(warehouse, "validate", "--type", "dromic", "--profile", "albay-pdrrmo-evacuation", "--file", str(ALBAY))
    assert scalar(db, "select count(*) from raw.dromic_evacuation") == before
    assert not list(warehouse["intake_dir"].glob("**/*.*"))


def test_missing_required_columns_is_a_file_error(warehouse, tmp_path):
    bad = tmp_path / "log.csv"
    bad.write_text("Request ID,Incident\nX-1,Carina\n", encoding="utf-8")
    report = run_intake(warehouse, "validate", "--type", "lgu", "--file", str(bad))
    assert not report["ok"]
    assert "Required columns not found" in report["file_errors"][0]


def test_layout_for_another_report_type_is_refused(warehouse):
    report = run_intake(warehouse, "validate", "--type", "dromic", "--profile", "quezon-city-cdrrmo-incidents",
                        "--file", str(QUEZON_CITY))
    assert not report["ok"] and "Response operations log" in report["file_errors"][0]


# ---------------------------------------------------------------------------- submission

@pytest.fixture
def submitted_quezon_city(warehouse):
    """Submit the Quezon City log into the e2e warehouse; afterwards remove it and rebuild."""
    report = run_intake(warehouse, "submit", "--type", "lgu", "--profile", "quezon-city-cdrrmo-incidents",
                        "--file", str(QUEZON_CITY))
    yield report
    shutil.rmtree(warehouse["intake_dir"], ignore_errors=True)
    subprocess.run([sys.executable, str(ROOT / "ingestion" / "pipeline.py"), "raw", "dbt"], cwd=ROOT,
                   env=warehouse["env"], capture_output=True, check=True)


def test_submission_loads_valid_rows_and_quarantines_the_rest(warehouse, submitted_quezon_city):
    report = submitted_quezon_city
    assert report["submitted"]["rows_written"] == 8          # the duplicate ticket is not stored
    assert report["submitted"]["rows_excluded"] == 1
    assert report["submitted"]["load"]["ok"], report["submitted"]["load"]   # every warehouse check passed
    with connect(warehouse["db"]) as conn:
        assert scalar(conn, "select count(*) from dw.fact_response_action where request_id like 'QC-2024-%'") == 5
        reasons = {r[0] for r in conn.execute(
            "select reject_reasons from dq.dq_quarantine q join meta.load_batch b on b.batch_id = "
            "(select max(batch_id) from meta.load_batch where source_file = q._source_file) where b.origin = 'intake'")}
        assert reasons == {"ARRIVAL_BEFORE_DISPATCH", "REQUEST_TYPE_UNKNOWN", "RESPONDING_UNIT_UNKNOWN"}
        # submitted files are real-world data: the synthetic-data proof is unaffected
        assert scalar(conn, "select count(*) from dq.dq_reconciliation where reconciliation_status <> 'match'") == 0


def test_resubmission_replaces_instead_of_duplicating(warehouse, submitted_quezon_city):
    again = run_intake(warehouse, "submit", "--type", "lgu", "--profile", "quezon-city-cdrrmo-incidents",
                       "--file", str(QUEZON_CITY))
    assert again["summary"] == submitted_quezon_city["summary"]   # its own IDs are not "already loaded"
    assert again["submitted"]["load"]["ok"]
    with connect(warehouse["db"]) as conn:
        assert scalar(conn, "select count(*) from dw.fact_response_action where request_id like 'QC-2024-%'") == 5


# ---------------------------------------------------------------------------- Excel export

def test_excel_export_has_every_mart_and_a_dictionary(warehouse, tmp_path):
    if "env" not in warehouse:
        pytest.skip("only run against the throwaway e2e database")
    out = tmp_path / "driw.xlsx"
    subprocess.run([sys.executable, str(ROOT / "ingestion" / "export_xlsx.py"), "--event", "TC2024-KRISTINE",
                    "--out", str(out)], cwd=ROOT, env=warehouse["env"], check=True, capture_output=True)
    wb = load_workbook(out)
    assert wb.sheetnames == ["About", "Event summary", "Impact cube", "Province impact", "Evacuation daily",
                             "Response performance", "Relief by organization", "Data dictionary"]
    summary = wb["Event summary"]
    assert summary.max_row == 2 and summary["A2"].value == "TC2024-KRISTINE"
    assert summary.freeze_panes == "A2" and summary.auto_filter.ref
    with connect(warehouse["db"]) as conn:
        assert wb["Evacuation daily"].max_row - 1 == scalar(
            conn, "select count(*) from mart.mart_evacuation_daily where event_code = 'TC2024-KRISTINE'")
    assert wb["Data dictionary"].max_row > 50
